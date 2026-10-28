# -*- coding: utf-8 -*-
"""
一人财税 · 经营报表生成器
--------------------------------
输入：三个月测试账单.csv
处理：
  1. 复用 tax_utils.process_bill / auto_classify 对全部收支条目自动分类（与 MCP 工具一致）
  2. 按月汇总收支，调用 calculate_tax 计算月度增值税、附加税、个人经营所得税
  3. 输出：
     - 三个月账单_分类明细.csv   分类后的明细账
     - 三个月经营报表.xlsx      多工作表汇总报表（带格式）
     - 三个月经营报表.md         可读版经营分析报告
"""
import sys

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from tax_utils import calculate_tax, process_bill

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

SRC_FILE = "三个月测试账单.csv"
DETAIL_CSV = "三个月账单_分类明细.csv"
REPORT_XLSX = "三个月经营报表.xlsx"
REPORT_MD = "三个月经营报表.md"

# ============ 1. 读取 + 自动分类 ============
df = process_bill(pd.read_csv(SRC_FILE))
df["交易日期"] = pd.to_datetime(df["交易日期"], errors="coerce")
df = df.dropna(subset=["交易日期"]).sort_values("交易日期").reset_index(drop=True)
df["月份"] = df["交易日期"].dt.strftime("%Y-%m")

detail = df[["月份", "交易日期", "交易描述", "交易金额", "大类", "明细分类"]].copy()
detail["交易日期"] = detail["交易日期"].dt.strftime("%Y-%m-%d")
detail.to_csv(DETAIL_CSV, index=False, encoding="utf-8-sig")

# ============ 2. 月度汇总 + 税额 ============
months = sorted(df["月份"].unique())
income_by_m = df[df["大类"] == "收入类"].groupby("月份")["交易金额"].sum()
cost_by_m = df[df["大类"] == "支出类"].groupby("月份")["交易金额"].sum()

summary_rows, tax_rows, tax_detail = [], [], {}
for m in months:
    income = float(income_by_m.get(m, 0))
    cost = float(cost_by_m.get(m, 0))
    profit = income - cost
    tax = calculate_tax(income, cost)
    tax_detail[m] = tax
    total_tax = tax["月度总税额"]
    summary_rows.append({
        "月份": m,
        "营业收入": round(income, 2),
        "成本支出": round(cost, 2),
        "净利润": round(profit, 2),
        "净利率(%)": round(profit / income * 100, 2) if income else 0.0,
        "成本率(%)": round(cost / income * 100, 2) if income else 0.0,
        "月度总税额": total_tax,
        "税后净利润": round(profit - total_tax, 2),
    })
    tax_rows.append({
        "月份": m,
        "营业收入": round(income, 2),
        "增值税": tax["增值税"],
        "附加税": tax["附加税"],
        "个人经营所得税(月估)": tax["个人经营所得税(月估)"],
        "月度总税额": total_tax,
        "综合税负率(%)": tax["综合税负率"],
    })

summary_df = pd.DataFrame(summary_rows)
tax_df = pd.DataFrame(tax_rows)

# ============ 3. 分类统计 ============
cat = (
    df.groupby(["大类", "明细分类"])["交易金额"]
    .agg(金额="sum", 笔数="count")
    .reset_index()
)
cat["金额"] = cat["金额"].round(2)
cat["占比(%)"] = cat.apply(
    lambda r: round(r["金额"] / (summary_df["营业收入"].sum() if r["大类"] == "收入类"
                                 else summary_df["成本支出"].sum()) * 100, 2), axis=1)
cat = cat.sort_values(["大类", "金额"], ascending=[True, False]).reset_index(drop=True)

# ============ 4. 总体指标 ============
tot_income = float(summary_df["营业收入"].sum())
tot_cost = float(summary_df["成本支出"].sum())
tot_profit = tot_income - tot_cost
tot_tax = float(tax_df["月度总税额"].sum())
overview = pd.DataFrame([
    ("统计期间", f"{months[0]} ~ {months[-1]}"),
    ("交易笔数", f"{len(df)} 笔"),
    ("营业收入合计", f"{tot_income:,.2f}"),
    ("成本支出合计", f"{tot_cost:,.2f}"),
    ("净利润", f"{tot_profit:,.2f}"),
    ("净利率", f"{tot_profit / tot_income * 100:.2f}%"),
    ("税额合计", f"{tot_tax:,.2f}"),
    ("税后净利润", f"{tot_profit - tot_tax:,.2f}"),
    ("综合税负率", f"{tot_tax / tot_income * 100:.2f}%"),
], columns=["指标", "数值"])

# ============ 5. 写 Excel ============
HEAD_FILL = PatternFill("solid", fgColor="2F5597")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
TOTAL_FILL = PatternFill("solid", fgColor="E8EEF7")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0.00'


def write_sheet(writer, name, data, money_cols=(), pct_cols=(), widths=None):
    data.to_excel(writer, sheet_name=name, index=False)
    ws = writer.sheets[name]
    for c, col in enumerate(data.columns, start=1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = (
            widths.get(col, 16) if widths else max(12, min(30, len(str(col)) * 2 + 6)))
        for r in range(1, len(data) + 2):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            if col in money_cols:
                cell.number_format = MONEY
            elif col in pct_cols:
                cell.number_format = '0.00'
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


with pd.ExcelWriter(REPORT_XLSX, engine="openpyxl") as writer:
    ov = overview.copy()
    ov.to_excel(writer, sheet_name="经营总览", index=False)
    ws = writer.sheets["经营总览"]
    for c in (1, 2):
        ws.cell(row=1, column=c).fill, ws.cell(row=1, column=c).font = HEAD_FILL, HEAD_FONT
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 26
    for r in range(1, len(ov) + 2):
        for c in (1, 2):
            ws.cell(row=r, column=c).border = BORDER

    write_sheet(writer, "月度收支汇总", summary_df,
                money_cols=("营业收入", "成本支出", "净利润", "月度总税额", "税后净利润"),
                pct_cols=("净利率(%)", "成本率(%)"), widths={"月份": 12})
    write_sheet(writer, "月度税额测算", tax_df,
                money_cols=("营业收入", "增值税", "附加税", "个人经营所得税(月估)", "月度总税额"),
                pct_cols=("综合税负率(%)",), widths={"月份": 12, "个人经营所得税(月估)": 22})
    write_sheet(writer, "收支分类统计", cat, money_cols=("金额",), pct_cols=("占比(%)",),
                widths={"明细分类": 16})
    write_sheet(writer, "分类明细", detail, money_cols=("交易金额",),
                widths={"交易日期": 14, "交易描述": 26, "明细分类": 16, "月份": 10})

    # 合计行
    ws = writer.sheets["月度收支汇总"]
    r = len(summary_df) + 2
    ws.cell(row=r, column=1, value="合计").font = Font(bold=True)
    for col, key in [(2, "营业收入"), (3, "成本支出"), (4, "净利润"),
                     (7, "月度总税额"), (8, "税后净利润")]:
        cell = ws.cell(row=r, column=col, value=round(float(summary_df[key].sum()), 2))
        cell.font, cell.number_format = Font(bold=True), MONEY
    for c in range(1, 9):
        ws.cell(row=r, column=c).fill = TOTAL_FILL
        ws.cell(row=r, column=c).border = BORDER

# ============ 6. 写 Markdown ============
top_costs = cat[cat["大类"] == "支出类"].head(3)
m_i = summary_df.set_index("月份")
g1 = (m_i["营业收入"].iloc[1] / m_i["营业收入"].iloc[0] - 1) * 100 if len(m_i) > 1 else 0
g2 = (m_i["营业收入"].iloc[2] / m_i["营业收入"].iloc[1] - 1) * 100 if len(m_i) > 2 else 0
c2 = (m_i["成本支出"].iloc[2] / m_i["成本支出"].iloc[1] - 1) * 100 if len(m_i) > 2 else 0

md = [
    f"# 一人财税 · 三个月经营报表（{months[0]} ~ {months[-1]}）",
    "",
    f"> 数据来源：`{SRC_FILE}`，共 {len(df)} 笔交易，已按关键词规则自动分类；"
    f"税额口径见文末「计税说明」。",
    "",
    "## 一、总体概况",
    "",
    "| 指标 | 数值 |",
    "| --- | ---: |",
]
md += [f"| {k} | {v} |" for k, v in overview.itertuples(index=False)]
md += ["", "## 二、月度收支汇总", "", summary_df.to_markdown(index=False), ""]
md += ["## 三、月度税额测算", "", tax_df.to_markdown(index=False), ""]
md += ["## 四、收支分类统计（三个月合计）", "", cat.to_markdown(index=False), ""]
md += [
    "## 五、经营洞察",
    "",
    f"1. **收入稳步增长**：{months[0]} 收入 {m_i['营业收入'].iloc[0]:,.0f} 元，"
    f"{months[1]} 环比 {g1:+.1f}%，{months[2]} 环比 {g2:+.1f}%，三个月累计 {tot_income:,.0f} 元。",
    f"2. **成本率 {tot_cost / tot_income * 100:.2f}%**：累计支出 {tot_cost:,.0f} 元，"
    f"净利润 {tot_profit:,.0f} 元，净利率 {tot_profit / tot_income * 100:.2f}%。",
    "3. **成本集中度**：前三大支出为 "
    + "、".join(f"{r.明细分类} {r.金额:,.0f} 元（{r.占比_pct:.1f}%）"
               for r in top_costs.rename(columns={"占比(%)": "占比_pct"}).itertuples())
    + "，合计占总支出的 "
    + f"{top_costs['金额'].sum() / tot_cost * 100:.1f}%。",
    f"4. **税负**：增值税因月销售额均低于 10 万元起征点全免，实际税负全部来自个人经营所得税，"
    f"累计 {tot_tax:,.0f} 元，综合税负率 {tot_tax / tot_income * 100:.2f}%；"
    f"税后净利润 {tot_profit - tot_tax:,.0f} 元。",
    f"5. **{months[1]} 支出异常**：当月支出 {m_i['成本支出'].iloc[1]:,.0f} 元，"
    f"较上月激增，主因进货采购 4,500 元、房租 2,500 元、外包 3,000 元同期发生；"
    f"{months[2]} 支出回落至 {m_i['成本支出'].iloc[2]:,.0f} 元（环比 {c2:+.1f}%）。",
    "",
    "## 六、计税说明与口径",
    "",
    "- **增值税**：小规模纳税人，月销售额 ≤ 10 万元免征；本期间各月收入 2.95 万 ~ 4.5 万元，均未达起征点，"
    "增值税为 0，附加税随之免征（若超过起征点，按 1% 征收率、附加税按增值税额 6% 计）。",
    "- **个人经营所得税**：按「月度利润 × 12」换算年度应纳税所得额，套用 5%~35% 五级超额累进税率，"
    "算出年度税额后除以 12 作为月度估算；未考虑专项扣除、社保及核定征收情形。",
    "- **成本确认**：云服务器年费、软件会员、域名续费按付款当期全额计入，未做跨期摊销；"
    "「退款收入-供应商」300 元按其他收入单列，未冲减进货成本。",
    "- **分类规则**：关键词匹配法，办公类优先于其他支出（如「办公水电费」计入办公成本），"
    "如需调整口径可修改 `tax_utils.CLASSIFY_RULES`。",
    "",
    "---",
    "",
    f"配套文件：分类明细 `{DETAIL_CSV}`；多工作表报表 `{REPORT_XLSX}`"
    "（经营总览 / 月度收支汇总 / 月度税额测算 / 收支分类统计 / 分类明细）。",
    "",
]

with open(REPORT_MD, "w", encoding="utf-8") as f:
    f.write("\n".join(md))

# ============ 7. 终端摘要 ============
print("=== 月度收支汇总 ===")
print(summary_df.to_string(index=False))
print("\n=== 月度税额测算 ===")
print(tax_df.to_string(index=False))
print("\n=== 收支分类统计 ===")
print(cat.to_string(index=False))
print(f"\n合计：收入 {tot_income:,.2f} | 支出 {tot_cost:,.2f} | 净利润 {tot_profit:,.2f} | 税额 {tot_tax:,.2f}")
print(f"已生成：{DETAIL_CSV} / {REPORT_MD} / {REPORT_XLSX}")
